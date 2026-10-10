"""Tests for the eval suite: scorers, pass@k math, the practice set's integrity, the runner, ablations.

Practice-set golds for generated files are re-derived here with code that is independent of
evals/practice/build_practice.py (openpyxl/csv/pypdf/json directly instead of the generator's
pandas paths), so a regenerated file that drifts from tasks.jsonl fails loudly.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import importlib.util
import json
import re
import shutil
import statistics
import subprocess
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from evals import ablate, passk  # noqa: E402
from evals import run_practice as rp  # noqa: E402
from evals.scorer import exact_match, gaia_match, question_scorer, score_answer  # noqa: E402

PRACTICE = ROOT / "evals" / "practice"
TASKS = [json.loads(line) for line in (PRACTICE / "tasks.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
BY_ID = {t["id"]: t for t in TASKS}
MANIFEST = json.loads((PRACTICE / "manifest.json").read_text(encoding="utf-8"))


def task_file(tid: str) -> Path:
    return ROOT / BY_ID[tid]["file"]


# =============================================================================== scorer


@pytest.mark.parametrize(
    "answer, gold, expected",
    [
        # numbers: only $ % , are stripped from the answer; units are fatal
        ("$1,234.50", "1234.5", True),
        ("1,000", "1000", True),
        ("17%", "17", True),
        ("17.0", "17", True),
        (" 17 ", "17", True),
        ("17 km", "17", False),
        ("seventeen", "17", False),
        ("17.01", "17", False),
        (None, "17", False),
        # lists: split on , and ; ; order and length matter; punctuation kept per element
        ("a, b; c", "a,b,c", True),
        ("b, a", "a, b", False),
        ("a, b", "a, b, c", False),
        ("1, 2.0", "1,2", True),
        ("$1, 2%", "1, 2", True),
        ("St. Louis, Paris", "St Louis, Paris", False),
        ("ST. LOUIS ,paris", "St. Louis, Paris", True),
        ("Argentina, Brazil, Chile, Paraguay, Peru", "Argentina, Brazil, Chile, Paraguay, Peru", True),
        # strings: whitespace + punctuation removed, case folded
        ("Sea gull.", "seagull", True),
        ("THE END!", "the end", True),
        ("Qe8", "Qe8#", True),
        ("Paris", "London", False),
        (None, "None", True),  # the reference maps a None answer to the string "None"
    ],
)
def test_gaia_scorer(answer, gold, expected):
    assert gaia_match(answer, gold) is expected


@pytest.mark.parametrize(
    "answer, gold, expected",
    [
        ("  Paris ", "paris", True),
        ("Paris.", "Paris", False),
        ("1,000", "1000", False),
        ("a,b", "a, b", False),
        ("17.0", "17", False),
        ("Qe8", "Qe8#", False),
        (None, "x", False),
        (17, "17", True),
    ],
)
def test_exact_match_is_the_course_rule(answer, gold, expected):
    assert exact_match(answer, gold) is expected


def test_score_answer_reports_both():
    assert score_answer("1,000", "1000") == {"gaia": True, "strict": False}
    assert score_answer("paris", "Paris") == {"gaia": True, "strict": True}


def test_question_scorer_is_faithful_on_reference_quirks():
    # unparseable numbers become inf, which never equals a finite ground truth
    assert question_scorer("n/a", "3") is False
    # a list ground truth with a numeric element compares that element numerically
    with pytest.warns(UserWarning):
        assert question_scorer("1, 2", "1, 2, 3") is False
    assert gaia_match(3.0, "3") is True  # wrapper coerces non-strings like an agent would submit
    assert gaia_match(["a", "b"], "a, b") is True


# =============================================================================== pass@k


def test_pass_at_k_hand_computed():
    assert passk.pass_at_k(5, 2, 2) == pytest.approx(0.7)  # 1 - C(3,2)/C(5,2) = 1 - 3/10
    assert passk.pass_at_k(10, 3, 5) == pytest.approx(1 - 21 / 252)  # C(7,5)=21, C(10,5)=252
    assert passk.pass_at_k(3, 1, 1) == pytest.approx(1 / 3)
    assert passk.pass_at_k(5, 0, 3) == 0.0
    assert passk.pass_at_k(5, 4, 2) == 1.0  # n - c < k
    assert passk.naive_pass_at_k(2 / 5, 2) == pytest.approx(0.64)  # the biased plug-in differs


def test_pass_hat_k_hand_computed():
    assert passk.pass_hat_k(5, 2, 2) == pytest.approx(0.1)  # C(2,2)/C(5,2) = 1/10
    assert passk.pass_hat_k(10, 3, 2) == pytest.approx(3 / 45)
    assert passk.pass_hat_k(5, 5, 3) == 1.0
    assert passk.pass_hat_k(5, 2, 3) == 0.0
    for n, c in [(4, 1), (6, 3), (7, 7)]:
        assert passk.pass_hat_k(n, c, 1) == pytest.approx(c / n) == pytest.approx(passk.pass_at_k(n, c, 1))


@pytest.mark.parametrize("n, c, k", [(3, 4, 1), (3, 1, 0), (3, 1, 4), (0, 0, 1)])
def test_passk_rejects_bad_inputs(n, c, k):
    with pytest.raises(ValueError):
        passk.pass_at_k(n, c, k)


def _doc(run_id, outcomes, cfg=None, strict=None):
    cfg = cfg or {"provider": "p", "model": "m", "arch": "single", "disabled_tools": [], "verify": False}
    items = []
    for tid, ok in outcomes.items():
        items.append({"id": tid, "answer": "x", "gold": "x", "correct": ok,
                      "correct_strict": ok if strict is None else strict.get(tid, ok),
                      "seconds": 1.0, "steps": 3, "input_tokens": 10, "output_tokens": 5, "error": None})
    return {"run_id": run_id, "config": cfg, "set": "practice", "items": items, "summary": {}}


def test_passk_analyze_over_runs():
    docs = [
        _doc("r1", {"A": True, "B": True, "C": False}),
        _doc("r2", {"A": True, "B": False, "C": False}),
        _doc("r3", {"A": True, "B": False, "C": False}),
    ]
    rep = passk.analyze(docs)
    rows = {r["k"]: r for r in rep["table"]}
    assert rows[1]["pass@k"] == pytest.approx((1 + 1 / 3 + 0) / 3, abs=1e-4)
    assert rows[3]["pass@k"] == pytest.approx((1 + 1 + 0) / 3, abs=1e-4)
    assert rows[3]["pass^k"] == pytest.approx(1 / 3, abs=1e-4)
    assert rep["flaky"] == ["B"] and rep["always_solved"] == ["A"] and rep["never_solved"] == ["C"]
    assert rep["same_config"] is True
    assert "pass^k" in passk.to_markdown(rep)


def test_passk_skips_skipped_items_and_flags_mixed_configs():
    d1 = _doc("r1", {"A": True})
    d2 = _doc("r2", {"A": False}, cfg={"provider": "p", "model": "other", "arch": "single", "disabled_tools": [], "verify": False})
    d2["items"].append({"id": "B", "skipped": True, "correct": None, "correct_strict": None})
    rep = passk.analyze([d1, d2])
    assert rep["per_task"] == {"A": {"n": 2, "c": 1, "p": 0.5}}
    assert rep["same_config"] is False


# =============================================================================== practice set integrity


def test_practice_set_shape():
    assert len(TASKS) == 16
    assert len(BY_ID) == 16, "ids must be unique"
    for t in TASKS:
        assert set(t) >= {"id", "question", "file", "gold", "format", "skills", "source"}, t["id"]
        assert t["format"] in {"number", "string", "list", "ordered_list"}
        assert t["question"].strip() and t["gold"].strip() and t["skills"]
    formats = {t["format"] for t in TASKS}
    assert {"number", "string", "ordered_list"} <= formats
    kinds = {s for t in TASKS for s in t["skills"] if s.startswith("file:")}
    assert kinds >= {"file:xlsx", "file:csv", "file:py", "file:audio", "file:image", "file:pdf", "file:json"}


def test_every_file_exists_and_matches_manifest():
    file_tasks = [t for t in TASKS if t["file"]]
    assert len(file_tasks) == 12
    for t in file_tasks:
        p = ROOT / t["file"]
        assert p.is_file() and p.stat().st_size > 0, t["file"]
        assert not p.resolve().is_relative_to((ROOT / "evals" / "results").resolve())
        entry = MANIFEST["files"][t["file"]]
        assert hashlib.sha256(p.read_bytes()).hexdigest() == entry["sha256"], f"{t['file']} changed; rebuild"


def test_golds_are_already_in_submission_form():
    from gaia_agent.agent import normalize_answer

    for t in TASKS:
        assert normalize_answer(t["gold"]) == t["gold"], t["id"]
        assert gaia_match(t["gold"], t["gold"]) and exact_match(t["gold"], t["gold"])


def test_wikipedia_tasks_are_sourced():
    wiki = [t for t in TASKS if t["source"]["type"] == "wikipedia"]
    assert len(wiki) == 4
    for t in wiki:
        assert t["file"] is None
        hops = t["source"]["hops"]
        assert hops and t["source"]["retrieved"] == "2026-10-10"
        for h in hops:
            assert h["url"].startswith("https://en.wikipedia.org/wiki/") and isinstance(h["revid"], int)
            assert h["retrieved"] == "2026-10-10" and len(h["evidence"]) > 20
        evidence = " ".join(h["evidence"] for h in hops)
        for part in t["gold"].split(", "):
            assert part in evidence, (t["id"], part)


# --- independent re-derivation of every generated-file gold ----------------------------------


def test_gold_p01_late_student_loans():
    from openpyxl import load_workbook

    ws = load_workbook(task_file("p01-xlsx-late-loans"), data_only=True)["Loans"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    n = 0
    for r in rows[1:]:
        rec = dict(zip(hdr, r))
        out, ret = rec["Checked Out"], rec["Returned"]
        if rec["Tier"] == "Student" and out.year == 2026 and out.month == 3 and ret is not None:
            n += (ret - out).days > 21
    assert str(n) == BY_ID["p01-xlsx-late-loans"]["gold"]


def test_gold_p02_combined_sku():
    from openpyxl import load_workbook

    wb = load_workbook(task_file("p02-xlsx-multisheet-sku"), data_only=True)
    assert len(wb.sheetnames) == 3
    totals: dict[str, int] = {}
    single_max = ("", -1)
    for ws in wb.worksheets:
        for sku, _desc, _bin, qty in list(ws.iter_rows(values_only=True))[1:]:
            totals[sku] = totals.get(sku, 0) + qty
            single_max = max(single_max, (sku, qty), key=lambda x: x[1])
    best = sorted(totals.items(), key=lambda kv: -kv[1])
    assert best[0][1] > best[1][1]
    assert best[0][0] == BY_ID["p02-xlsx-multisheet-sku"]["gold"]
    assert single_max[0] != best[0][0], "trap: the largest single row is a different SKU"


def test_gold_p03_evening_long_trips():
    n = 0
    with open(task_file("p03-csv-evening-trips"), newline="") as fh:
        for r in csv.DictReader(fh):
            dep = dt.datetime.strptime(r["departure"], "%Y-%m-%d %H:%M")
            arr = dt.datetime.strptime(r["arrival"], "%Y-%m-%d %H:%M")
            n += dep.hour >= 18 and (arr - dep) > dt.timedelta(minutes=90)
    assert str(n) == BY_ID["p03-csv-evening-trips"]["gold"]


def test_gold_p04_pm25_ranking_ignores_blanks():
    vals: dict[str, list[float]] = {}
    with open(task_file("p04-csv-pm25-ranking"), newline="") as fh:
        for r in csv.DictReader(fh):
            vals.setdefault(r["station"], [])
            if r["pm25_ugm3"].strip():
                vals[r["station"]].append(float(r["pm25_ugm3"]))
    means = sorted(((statistics.mean(v), s) for s, v in vals.items()), reverse=True)
    assert ", ".join(s for _, s in means[:3]) == BY_ID["p04-csv-pm25-ranking"]["gold"]


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_gold_p05_queue_sim():
    t = BY_ID["p05-py-queue-sim"]
    assert "240 minutes" in t["question"] and "three counters" in t["question"]
    mod = _load(task_file("p05-py-queue-sim"), "t_queue_sim")
    assert str(mod.simulate(240, counters=3)) == t["gold"]
    assert str(mod.simulate(60)) != t["gold"], "the default printout must not be the answer"


def test_gold_p06_cipher_roundtrip():
    t = BY_ID["p06-py-invert-cipher"]
    target = re.search(r"encode\(w\) == '([a-z]+)'", t["question"]).group(1)
    mod = _load(task_file("p06-py-invert-cipher"), "t_cipher")
    assert t["gold"].isalpha() and t["gold"].islower()
    assert mod.encode(t["gold"]) == target


WORDS = {"four": 4, "six": 6, "nine": 9, "eleven": 11}


def test_gold_p07_from_tts_script():
    t = BY_ID["p07-audio-platform"]
    tts = t["source"]["tts_text"]
    sentence = next(s for s in tts.split(". ") if s.startswith("Passengers for Preston"))
    word = re.search(r"platform (\w+)", sentence).group(1)
    assert str(WORDS[word]) == t["gold"]
    assert "Preston is cancelled" in tts, "trap: the first Preston service is cancelled"


def test_gold_p08_from_tts_script():
    t = BY_ID["p08-audio-mean-reading"]
    readings = [Decimal(x) for x in re.findall(r"\d+\.\d", t["source"]["tts_text"])]
    assert len(readings) == 5
    mean = sum(readings) / len(readings)
    assert str(mean.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)) == t["gold"]


def test_gold_p09_unique_mate_and_image_matches_fen(tmp_path):
    import chess

    t = BY_ID["p09-image-chess-mate"]
    board = chess.Board(t["source"]["fen"])
    assert board.turn == chess.WHITE and board.is_valid()
    mates = []
    for m in list(board.legal_moves):
        san = board.san(m)
        board.push(m)
        if board.is_checkmate():
            mates.append(san)
        board.pop()
    assert mates == [t["gold"]]
    _assert_rerender_identical(tmp_path, "build_p09", task_file("p09-image-chess-mate"))


def test_gold_p10_chart_values_and_render(tmp_path):
    t = BY_ID["p10-image-chart-mean"]
    values = t["source"]["values"]
    mean = sum(values.values()) / len(values)
    ranked = sorted(values, key=lambda m: abs(values[m] - mean))
    assert ranked[0] == t["gold"]
    assert abs(values[ranked[1]] - mean) - abs(values[ranked[0]] - mean) >= 1
    _assert_rerender_identical(tmp_path, "build_p10", task_file("p10-image-chart-mean"))


def _assert_rerender_identical(tmp_path: Path, builder: str, committed: Path) -> None:
    """The committed PNG must be exactly what the stored data renders (ties image to gold)."""
    from PIL import Image, ImageChops

    from evals.practice import build_practice as bp

    try:
        out = tmp_path / committed.name
        getattr(bp, builder)(out)
    except RuntimeError as e:  # no usable font on this machine
        pytest.skip(str(e))
    a, b = Image.open(committed).convert("RGB"), Image.open(out).convert("RGB")
    if a.size != b.size or ImageChops.difference(a, b).getbbox() is not None:
        pytest.skip("font rasterisation differs on this machine; cannot pixel-compare")


def test_gold_p11_pdf_action_items():
    from pypdf import PdfReader

    text = "\n".join(p.extract_text() or "" for p in PdfReader(str(task_file("p11-pdf-action-items"))).pages)
    n = 0
    for line in text.splitlines():
        m = re.match(r"AI-\d+ \| .* \| Owner: (.+) \| Due: (\d{4})-(\d{2})-(\d{2})$", line.strip())
        if m and m.group(1) == "Priya Raman":
            n += dt.date(int(m.group(2)), int(m.group(3)), int(m.group(4))) < dt.date(2026, 11, 15)
    assert str(n) == BY_ID["p11-pdf-action-items"]["gold"]


def test_gold_p12_top_shipped_customer():
    doc = json.loads(task_file("p12-json-top-customer").read_text())
    spend: dict[str, float] = {}
    for o in doc["orders"]:
        if o["status"] != "shipped":
            continue
        name = o["customer"]["first"] + " " + o["customer"]["last"]
        spend[name] = spend.get(name, 0.0) + sum(i["qty"] * i["unit_price"] for i in o["items"])
    assert max(spend, key=spend.get) == BY_ID["p12-json-top-customer"]["gold"]


def test_audio_files_are_real_audio():
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        pytest.skip("ffprobe not installed")
    for tid in ("p07-audio-platform", "p08-audio-mean-reading"):
        out = subprocess.run([ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                              str(task_file(tid))], capture_output=True, text=True, check=True).stdout
        assert 5 < float(out.strip()) < 60, (tid, out)


@pytest.mark.network
def test_audio_transcribes_to_its_script():
    """Whisper (base) hears what the TTS script says. May download the model on first run."""
    np = pytest.importorskip("numpy")
    fw = pytest.importorskip("faster_whisper")
    if not shutil.which("ffmpeg"):
        pytest.skip("ffmpeg not installed")
    model = fw.WhisperModel("base", device="cpu", compute_type="int8")
    for tid, needles in (("p07-audio-platform", ["Preston", "11"]), ("p08-audio-mean-reading", ["21.4", "18.9"])):
        raw = subprocess.run(["ffmpeg", "-loglevel", "error", "-i", str(task_file(tid)), "-f", "s16le", "-ac", "1",
                              "-ar", "16000", "-"], capture_output=True, check=True).stdout
        audio = np.frombuffer(raw, np.int16).astype(np.float32) / 32768
        text = " ".join(s.text for s in model.transcribe(audio, language="en")[0])
        for needle in needles:
            assert needle in text, (tid, needle, text)


@pytest.mark.network
def test_practice_questions_do_not_paraphrase_course_questions():
    """Contamination guard: no practice question is a near-copy of a course question (questions only)."""
    requests = pytest.importorskip("requests")
    try:
        qs = requests.get("https://agents-course-unit4-scoring.hf.space/questions", timeout=30).json()
    except Exception as e:
        pytest.skip(f"course API unreachable: {e}")

    def toks(s):
        return set(re.findall(r"[a-z0-9]+", s.lower())) - {"the", "a", "an", "of", "in", "and", "to", "is", "what", "which"}

    for t in TASKS:
        for q in qs:
            a, b = toks(t["question"]), toks(q["question"])
            assert len(a & b) / max(1, len(a | b)) < 0.35, (t["id"], q["task_id"])


# =============================================================================== runner


def _stub(answers: dict[str, str], steps: int = 4, calls: list | None = None, error: dict | None = None):
    def solve(question, task_id, file_path, max_steps, stop_event):
        if calls is not None:
            calls.append({"id": task_id, "file_path": file_path, "max_steps": max_steps, "question": question})
        if error and task_id in error:
            return {"answer": "", "error": error[task_id], "steps": 1, "seconds": 0.1}
        return {"answer": answers.get(task_id, "wrong"), "steps": steps, "seconds": 0.5, "error": None,
                "input_tokens": 100, "output_tokens": 20}
    return solve


def _run_id(tag):
    return f"pytest-evals-{tag}-{dt.datetime.now().strftime('%H%M%S%f')}"


def test_run_writes_schema_and_scores_both_ways(tmp_path):
    tasks = rp.select_tasks(rp.load_tasks(), "p01,p02,p04,w03")
    answers = {"p01-xlsx-late-loans": BY_ID["p01-xlsx-late-loans"]["gold"] + ".0",  # gaia yes, strict no
               "p02-xlsx-multisheet-sku": BY_ID["p02-xlsx-multisheet-sku"]["gold"].lower(),  # both yes
               "w03-wiki-bolivia-borders": "Argentina,Brazil,Chile,Paraguay,Peru"}  # gaia yes, strict no
    calls: list = []
    doc = rp.run(tasks, run_id=_run_id("schema"), solve_fn=_stub(answers, calls=calls), results_dir=tmp_path,
                 staging=tmp_path / "staging", config={"model": "stub", "arch": "single", "disabled_tools": [],
                                                       "verify": False}, log=None)
    saved = json.loads((tmp_path / f"{doc['run_id']}.json").read_text())
    assert saved["set"] == "practice" and set(saved["config"]) >= {"model", "arch", "disabled_tools", "verify"}
    for it in saved["items"]:
        assert set(it) >= {"id", "answer", "gold", "correct", "seconds", "steps", "input_tokens", "output_tokens", "error"}
    s = saved["summary"]
    assert (s["n"], s["correct"], s["correct_strict"], s["format_gap"]) == (4, 3, 1, 2)
    assert s["accuracy"] == 0.75 and s["total_tokens"] == 4 * 120 and s["mean_steps"] == 4
    assert s["by_format"]["ordered_list"] == {"n": 2, "correct": 1}
    lo, hi = s["accuracy_ci95"]
    assert lo < 0.75 < hi
    staged = [c["file_path"] for c in calls if c["file_path"]]
    assert staged and all(Path(p).is_file() and "evals" not in Path(p).relative_to(tmp_path).parts for p in staged)
    assert next(c for c in calls if c["id"] == "w03-wiki-bolivia-borders")["file_path"] is None


def test_budget_guard_stops_and_marks_skipped(tmp_path):
    tasks = rp.select_tasks(rp.load_tasks(), limit=5)
    calls: list = []
    doc = rp.run(tasks, run_id=_run_id("budget"), solve_fn=_stub({}, steps=5, calls=calls), max_requests=12,
                 results_dir=tmp_path, staging=tmp_path / "s", config={}, log=None)
    assert len(calls) == 2  # 5 + 5 calls used; 2 left < MIN_TASK_BUDGET
    assert calls[0]["max_steps"] == 11 and calls[1]["max_steps"] == 6  # capped by what is left
    skipped = [it for it in doc["items"] if it.get("skipped")]
    assert len(skipped) == 3 and doc["summary"]["skipped"] == 3 and doc["summary"]["n"] == 2
    assert "budget" in doc["budget"]["stopped"]


def test_provider_quota_error_stops_run_without_scoring(tmp_path):
    tasks = rp.select_tasks(rp.load_tasks(), limit=3)
    err = {tasks[1]["id"]: "AgentGenerationError: Error in generating model output:\nDaily request cap reached for 'openrouter'"}
    doc = rp.run(tasks, run_id=_run_id("quota"), solve_fn=_stub({}, error=err), results_dir=tmp_path,
                 staging=tmp_path / "s", config={}, log=None)
    assert [bool(it.get("skipped")) for it in doc["items"]] == [False, True, True]
    assert doc["summary"]["n"] == 1 and doc["budget"]["stopped"].startswith("provider quota")


def test_resume_keeps_scored_items(tmp_path):
    tasks = rp.select_tasks(rp.load_tasks(), limit=3)
    rid = _run_id("resume")
    rp.run(tasks, run_id=rid, solve_fn=_stub({}), max_requests=6, results_dir=tmp_path, staging=tmp_path / "s",
           config={}, log=None)
    calls: list = []
    doc = rp.run(tasks, run_id=rid, solve_fn=_stub({}, calls=calls), resume=True, results_dir=tmp_path,
                 staging=tmp_path / "s", config={}, log=None)
    assert [c["id"] for c in calls] == [t["id"] for t in tasks[1:]]
    assert doc["summary"]["n"] == 3 and doc["summary"]["skipped"] == 0


def test_select_tasks_by_prefix_and_limit():
    tasks = rp.load_tasks()
    assert [t["id"] for t in rp.select_tasks(tasks, "w")] == [t["id"] for t in TASKS if t["id"].startswith("w")]
    assert len(rp.select_tasks(tasks, "p0", limit=3)) == 3


def test_default_solve_attaches_local_file(monkeypatch):
    import gaia_agent.agent as agent_mod

    seen = {}

    def fake_solve(question, task_id=None, file_name=None, *, agent=None, stop_event=None, max_steps=12):
        seen.update(question=question, task_id=task_id, file_name=file_name, max_steps=max_steps)
        return {"answer": "x"}

    monkeypatch.setattr(agent_mod, "solve", fake_solve)
    rp.default_solve("Q?", "p01", "/abs/file.xlsx", 7, None)
    assert seen["question"].endswith("Attached file: /abs/file.xlsx")
    assert seen["file_name"] is None and seen["max_steps"] == 7


def test_wilson_interval_hand_computed():
    lo, hi = rp.wilson_ci(8, 16)
    assert lo == pytest.approx(0.2800, abs=1e-3) and hi == pytest.approx(0.7200, abs=1e-3)
    assert rp.wilson_ci(0, 0) is None
    assert rp.wilson_ci(0, 16)[0] == 0.0


# =============================================================================== ablations


def test_configs_are_valid_and_name_real_tools():
    cfgs = ablate.load_configs()
    assert {"baseline", "no_web", "no_media", "no_wikipedia", "arch_multi", "verify"} <= set(cfgs)
    assert any(c["group"] == "model" for c in cfgs.values())
    tools = ablate.known_tool_names()
    if tools is None:
        pytest.skip("agent tool modules not importable here")
    for c in cfgs.values():
        assert ablate.validate_config(c, tools) == [], c["name"]
    assert ablate.validate_config({"name": "x", "env": {"GAIA_DISABLE_TOOLS": "not_a_tool"}}, tools)


def test_select_configs_always_includes_baseline():
    cfgs = ablate.load_configs()
    names = [c["name"] for c in ablate.select_configs(cfgs, "no_web")]
    assert names == ["baseline", "no_web"]
    assert all(c["default"] for c in ablate.select_configs(cfgs))


def test_mcnemar_exact_hand_computed():
    assert ablate.mcnemar_exact(0, 0) == 1.0
    assert ablate.mcnemar_exact(5, 0) == pytest.approx(2 / 32)
    assert ablate.mcnemar_exact(6, 1) == pytest.approx(2 * 8 / 128)
    assert ablate.mcnemar_exact(3, 3) == 1.0


def test_aggregate_builds_comparison(tmp_path):
    tag = "t1"
    base = _doc(f"abl-{tag}-baseline", {"p01-xlsx-late-loans": True, "w01-wiki-element-city": True,
                                       "w02-wiki-architect-prize": True})
    noweb = _doc(f"abl-{tag}-no_web", {"p01-xlsx-late-loans": True, "w01-wiki-element-city": False,
                                     "w02-wiki-architect-prize": False})
    for d in (base, noweb):
        (tmp_path / f"{d['run_id']}.json").write_text(json.dumps(d))
    cfgs = [{"name": "baseline", "env": {}}, {"name": "no_web", "env": {"GAIA_DISABLE_TOOLS": "web_search"}},
            {"name": "verify", "env": {"GAIA_VERIFY": "1"}}]
    rep = ablate.aggregate(cfgs, tag, tmp_path, [BY_ID[k] for k in ("p01-xlsx-late-loans", "w01-wiki-element-city",
                                                                    "w02-wiki-architect-prize")])
    rows = {r["name"]: r for r in rep["configs"]}
    assert rows["baseline"]["accuracy"] == 1.0
    vb = rows["no_web"]["vs_baseline"]
    assert vb["delta"] == pytest.approx(-2 / 3, abs=1e-3) and (vb["baseline_only"], vb["config_only"]) == (2, 0)
    assert vb["mcnemar_p"] == 0.5
    assert rows["verify"]["status"] == "not run"
    assert rep["per_skill"]["wikipedia"] == {"baseline": "2/2", "no_web": "0/2"}
    md = ablate.to_markdown(rep)
    assert "| no_web |" in md and "-66.7%" in md


# =============================================================================== end-to-end (no API)


def test_end_to_end_with_scripted_model(tmp_path, monkeypatch):
    """Runner -> real gaia_agent.agent.solve -> real CodeAgent + read_file tool -> scorer, with a
    scripted offline model standing in for the LLM. Zero API calls; proves the plumbing, not the model."""
    agent_mod = pytest.importorskip("gaia_agent.agent")
    from smolagents.models import ChatMessage, MessageRole, Model
    from smolagents.monitoring import TokenUsage

    if not hasattr(agent_mod, "get_model"):
        pytest.skip("gaia_agent.agent no longer builds its model through get_model()")

    class Scripted(Model):
        def generate(self, messages, stop_sequences=None, response_format=None, tools_to_call_from=None, **kw):
            path = re.search(r"Attached file: (\S+?\.pdf)", json.dumps([str(m) for m in messages])).group(1)
            code = (
                "import re\n"
                f"text = read_file({path!r})\n"
                "rows = re.findall(r'Owner: (.+?) \\| Due: (\\d{4}-\\d{2}-\\d{2})', text)\n"
                "n = sum(1 for o, d in rows if o.strip() == 'Priya Raman' and d < '2026-11-15')\n"
                "final_answer(str(n))"
            )
            return ChatMessage(role=MessageRole.ASSISTANT, content=f"Thought: count the rows.\n<code>\n{code}\n</code>",
                               token_usage=TokenUsage(input_tokens=50, output_tokens=10))

    monkeypatch.setattr(agent_mod, "get_model", lambda *a, **k: Scripted(model_id="scripted"))
    monkeypatch.delenv("GAIA_ARCH", raising=False)
    monkeypatch.delenv("GAIA_VERIFY", raising=False)
    tasks = rp.select_tasks(rp.load_tasks(), "p11")
    doc = rp.run(tasks, run_id=_run_id("e2e"), results_dir=tmp_path, staging=tmp_path / "staging",
                 config={"model": "scripted", "arch": "single", "disabled_tools": [], "verify": False},
                 timeout=120, log=None)
    item = doc["items"][0]
    assert item["error"] is None, item
    assert item["answer"] == BY_ID["p11-pdf-action-items"]["gold"]
    assert item["correct"] and item["correct_strict"] and item["steps"] >= 1
    # wave-3 wiring: the real solve() wrote a full trace under the run id, and the item points at it
    trace = json.loads(Path(item["trace_path"]).read_text(encoding="utf-8"))
    assert trace["run_id"] == doc["run_id"] and trace["task_id"] == "p11-pdf-action-items"
    assert trace["arch"] == "single" and trace["totals"]["tool_calls_by_name"].get("read_file") == 1
    assert item["model_calls"] == trace["totals"]["model_calls"] and item["input_tokens"] == 50 * item["model_calls"]
