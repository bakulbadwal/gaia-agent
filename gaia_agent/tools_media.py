"""Media and file tools for the GAIA agent (smolagents @tool functions).

Every tool returns a string and never raises: failures come back as text that
starts with "ERROR:", so the agent can read the problem and try something else.

Exports TOOLS_MEDIA.
"""

from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import zipfile
from pathlib import Path

from smolagents import tool

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MAX_CHARS = 40_000

TEXT_EXTS = {
    ".txt", ".md", ".py", ".json", ".jsonl", ".xml", ".html", ".htm", ".yaml", ".yml",
    ".tsv", ".log", ".ini", ".cfg", ".toml", ".js", ".ts", ".java", ".c", ".cpp", ".h",
    ".rs", ".go", ".rb", ".sh", ".sql", ".tex", ".bib", ".srt", ".vtt", ".pdb",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac", ".aiff", ".opus"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}


# --------------------------------------------------------------------------- helpers


def _data_dir() -> Path:
    """DATA_DIR from gaia_agent.config if it imports, else <project>/data."""
    try:
        from gaia_agent.config import DATA_DIR  # noqa: WPS433 (lazy on purpose)

        return Path(DATA_DIR)
    except Exception:
        return PROJECT_ROOT / "data"


def _ffmpeg() -> str:
    return shutil.which("ffmpeg") or "/opt/homebrew/bin/ffmpeg"


def _ffprobe() -> str:
    return shutil.which("ffprobe") or "/opt/homebrew/bin/ffprobe"


def _python_exe() -> str:
    venv_py = PROJECT_ROOT / ".venv" / "bin" / "python"
    return str(venv_py) if venv_py.exists() else sys.executable


def _truncate(text: str, limit: int = MAX_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n... [truncated: showing {limit} of {len(text)} characters]"


def _resolve(path: str) -> Path | None:
    """Find a file: as given, then under data/files/, then data/files/*_<name>."""
    if not path:
        return None
    p = Path(os.path.expanduser(str(path).strip().strip("'\"")))
    if p.exists():
        return p
    files_dir = _data_dir() / "files"
    cand = files_dir / p.name
    if cand.exists():
        return cand
    if files_dir.exists():
        matches = sorted(files_dir.glob(f"*_{p.name}")) or sorted(files_dir.glob(f"{p.name}*"))
        if matches:
            return matches[0]
    return None


def _not_found(path: str) -> str:
    return (
        f"ERROR: file not found: {path!r}. Use the exact local path returned when the task "
        "file was downloaded (files live under data/files/)."
    )


def _fmt_ts(seconds: float) -> str:
    seconds = max(0, int(seconds or 0))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _df_markdown(df) -> str:
    try:
        return df.to_markdown(index=False)
    except Exception:
        return df.to_string(index=False)


def _describe_dataframe(df, name: str, max_rows: int = 300) -> str:
    out = [f"### {name}", f"rows: {len(df)}, columns: {len(df.columns)}"]
    out.append("column dtypes: " + ", ".join(f"{c} ({t})" for c, t in df.dtypes.astype(str).items()))
    # Columns that look like categories / types (few distinct text values).
    cats = []
    for col in df.columns:
        s = df[col]
        from pandas.api import types as pdt

        if pdt.is_object_dtype(s) or pdt.is_string_dtype(s) or pdt.is_bool_dtype(s) or isinstance(
            s.dtype, __import__("pandas").CategoricalDtype
        ):
            vals = s.dropna().astype(str).unique()
            if 0 < len(vals) <= 25 and (len(vals) < max(2, len(df) * 0.6) or len(df) <= 5):
                cats.append(f"  - {col}: {sorted(vals)[:25]}")
    if cats:
        out.append("category-like columns (distinct values):\n" + "\n".join(cats))
    num = df.select_dtypes("number")
    if not num.empty and len(df) > 1:
        sums = ", ".join(f"{c}={num[c].sum():g}" for c in num.columns[:30])
        out.append(f"numeric column sums: {sums}")
    if len(df) > max_rows:
        out.append(f"(showing first {max_rows} of {len(df)} rows)")
        out.append(_df_markdown(df.head(max_rows)))
    else:
        out.append(_df_markdown(df))
    return "\n".join(out)


def _excel_fill_colors(path: Path, limit: int = 400) -> str:
    """List cells with a solid background fill, grouped by colour (openpyxl, .xlsx only)."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return ""
    parts = []
    for ws in wb.worksheets:
        by_color: dict[str, list[str]] = {}
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                fill = getattr(cell, "fill", None)
                if fill is None or fill.fill_type != "solid":
                    continue
                fg = fill.fgColor
                if fg is None:
                    continue
                if fg.type == "rgb" and isinstance(fg.rgb, str):
                    color = fg.rgb
                elif fg.type == "theme":
                    color = f"theme{fg.theme}" + (f"(tint {fg.tint:+.2f})" if fg.tint else "")
                elif fg.type == "indexed":
                    color = f"indexed{fg.indexed}"
                else:
                    continue
                if color in ("00000000",):
                    continue
                val = "" if cell.value is None else f"={cell.value!r}"
                by_color.setdefault(color, []).append(f"{cell.coordinate}{val}")
                count += 1
        if by_color:
            lines = [f"### cell fill colours in sheet '{ws.title}' ({count} filled cells)"]
            shown = 0
            for color, cells in sorted(by_color.items(), key=lambda kv: -len(kv[1])):
                take = cells[: max(0, limit - shown)]
                shown += len(take)
                more = f" ... (+{len(cells) - len(take)} more)" if len(take) < len(cells) else ""
                lines.append(f"- {color} ({len(cells)} cells): " + ", ".join(take) + more)
            parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _read_excel(p: Path) -> str:
    import pandas as pd

    engine = "xlrd" if p.suffix.lower() == ".xls" else None
    sheets = pd.read_excel(p, sheet_name=None, engine=engine)
    out = [f"Excel file {p.name}: {len(sheets)} sheet(s): {list(sheets)}"]
    for name, df in sheets.items():
        out.append(_describe_dataframe(df, f"sheet '{name}'"))
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        colors = _excel_fill_colors(p)
        if colors:
            out.append(colors)
    return "\n\n".join(out)


def _read_csv(p: Path) -> str:
    import pandas as pd

    sep = "\t" if p.suffix.lower() == ".tsv" else None
    df = pd.read_csv(p, sep=sep, engine="python")
    return f"CSV file {p.name}\n" + _describe_dataframe(df, p.name)


def _read_pdf(p: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "ERROR: no PDF library installed (pip install pypdf)."
    reader = PdfReader(str(p))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"--- page {i} ---\n{(page.extract_text() or '').strip()}")
    return f"PDF {p.name}: {len(reader.pages)} page(s)\n" + "\n".join(pages)


def _read_docx(p: Path) -> str:
    import docx

    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"--- table {ti} ---")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return f"Word document {p.name}\n" + "\n".join(parts)


def _read_pptx(p: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(p))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        parts.append(f"--- slide {i} ---\n" + "\n".join(t for t in texts if t.strip()))
    return f"PowerPoint {p.name}: {len(prs.slides)} slide(s)\n" + "\n".join(parts)


def _read_zip(p: Path) -> str:
    dest = p.parent / f"{p.stem}_unzipped"
    dest.mkdir(parents=True, exist_ok=True)
    names = []
    with zipfile.ZipFile(p) as zf:
        for member in zf.infolist():
            target = (dest / member.filename).resolve()
            if not str(target).startswith(str(dest.resolve())):
                continue  # skip path traversal entries
            if member.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(member) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)
            names.append(f"{target} ({member.file_size} bytes)")
    return (
        f"ZIP {p.name} extracted to {dest} ({len(names)} files). Call read_file on each path:\n"
        + "\n".join(names)
    )


def _read_text(p: Path) -> str:
    raw = p.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if p.suffix.lower() == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
    return f"File {p.name} ({len(raw)} bytes):\n{text}"


def _read_file_impl(path: str) -> str:
    p = _resolve(path)
    if p is None:
        return _not_found(path)
    if p.is_dir():
        entries = sorted(str(x) for x in p.iterdir())
        return f"{p} is a directory with {len(entries)} entries:\n" + "\n".join(entries[:200])
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _read_excel(p)
    if ext in (".csv", ".tsv"):
        try:
            return _read_csv(p)
        except Exception as e:  # malformed CSV: fall back to raw text
            return f"(could not parse as a table: {e})\n" + _read_text(p)
    if ext == ".pdf":
        return _read_pdf(p)
    if ext == ".docx":
        return _read_docx(p)
    if ext == ".pptx":
        return _read_pptx(p)
    if ext == ".zip":
        return _read_zip(p)
    if ext in IMAGE_EXTS:
        size = ""
        try:
            from PIL import Image

            with Image.open(p) as im:
                size = f" ({im.width}x{im.height} px, mode {im.mode})"
        except Exception:
            pass
        return (
            f"{p} is an image{size}. Use image_question(path={str(p)!r}, question=...) "
            "to ask the vision model about it."
        )
    if ext in AUDIO_EXTS:
        return f"{p} is an audio file. Use transcribe_audio(path={str(p)!r}) to get its transcript."
    if ext in VIDEO_EXTS:
        return (
            f"{p} is a video file. Use transcribe_audio(path={str(p)!r}) for its speech; "
            "frame sampling (youtube_frames_question) only works on YouTube URLs."
        )
    if ext in TEXT_EXTS or ext == "":
        return _read_text(p)
    # Unknown extension: try as text, else report binary.
    raw = p.read_bytes()[:4096]
    if b"\x00" in raw:
        return f"{p} is a binary file ({p.stat().st_size} bytes, extension {ext!r}); no reader for it."
    return _read_text(p)


# ----------------------------------------------------------------- whisper helpers

_WHISPER_MODELS: dict[str, object] = {}
_WHISPER_LOCK = threading.Lock()


def _whisper_model_name() -> str:
    return os.environ.get("GAIA_WHISPER_MODEL", "small")


def _get_whisper(name: str):
    with _WHISPER_LOCK:
        if name not in _WHISPER_MODELS:
            from faster_whisper import WhisperModel

            _WHISPER_MODELS[name] = WhisperModel(name, device="cpu", compute_type="int8")
        return _WHISPER_MODELS[name]


def _file_hash(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:20]


def _decode_audio(p: Path):
    """Decode any audio/video file to 16 kHz mono float32 with ffmpeg.

    faster-whisper's own PyAV decoder breaks on newer PyAV releases (it passes the removed
    `metadata_errors` argument), so ffmpeg is the primary path; the file path is the fallback.
    """
    import numpy as np

    ffmpeg = _ffmpeg()
    if not Path(ffmpeg).exists():
        return str(p)
    r = subprocess.run(
        [ffmpeg, "-nostdin", "-loglevel", "error", "-i", str(p), "-vn", "-f", "s16le",
         "-acodec", "pcm_s16le", "-ac", "1", "-ar", "16000", "-"],
        capture_output=True, timeout=1800,
    )
    if r.returncode != 0 or not r.stdout:
        raise RuntimeError(f"ffmpeg could not decode {p.name}: {r.stderr.decode(errors='replace')[-500:]}")
    return np.frombuffer(r.stdout, dtype=np.int16).astype(np.float32) / 32768.0


def _render_segments(segments: list[dict], timestamps: bool) -> str:
    if timestamps:
        return "\n".join(f"[{_fmt_ts(s['start'])}-{_fmt_ts(s['end'])}] {s['text']}" for s in segments)
    return " ".join(s["text"] for s in segments).strip()


def _transcribe(p: Path, timestamps: bool = False, cache_key: str | None = None) -> str:
    """Transcribe a local audio/video file with faster-whisper, caching by content hash."""
    model_name = _whisper_model_name()
    tdir = _data_dir() / "transcripts"
    tdir.mkdir(parents=True, exist_ok=True)
    key = cache_key or _file_hash(p)
    stem = f"{key}_{model_name.replace('/', '-')}"
    txt_path, seg_path = tdir / f"{stem}.txt", tdir / f"{stem}.segments.json"
    if seg_path.exists():
        data = json.loads(seg_path.read_text())
        segments, lang = data["segments"], data.get("language", "?")
    else:
        model = _get_whisper(model_name)
        seg_iter, info = model.transcribe(_decode_audio(p), beam_size=5, vad_filter=False)
        segments = [{"start": s.start, "end": s.end, "text": s.text.strip()} for s in seg_iter]
        lang = getattr(info, "language", "?")
        seg_path.write_text(json.dumps({"language": lang, "model": model_name, "segments": segments}))
        txt_path.write_text(_render_segments(segments, False))
    if not segments:
        return f"(no speech detected in {p.name})"
    header = f"Transcript of {p.name} (whisper {model_name}, language {lang}):\n"
    return _truncate(header + _render_segments(segments, timestamps))


# ----------------------------------------------------------------- youtube helpers


def _video_id(url: str) -> str:
    m = re.search(r"(?:v=|youtu\.be/|shorts/|embed/|live/)([A-Za-z0-9_-]{11})", url)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", url.strip()):
        return url.strip()
    return hashlib.sha256(url.encode()).hexdigest()[:16]


def _normalize_yt_url(url: str) -> str:
    url = url.strip()
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", url):
        return f"https://www.youtube.com/watch?v={url}"
    return url


def _ydl(**opts):
    import yt_dlp

    base = {
        "quiet": True,
        "no_warnings": True,
        "noprogress": True,
        "noplaylist": True,
        "socket_timeout": 30,
        "retries": 3,
        "ffmpeg_location": str(Path(_ffmpeg()).parent),
    }
    base.update(opts)
    return yt_dlp.YoutubeDL(base)


def _pick_track(tracks: dict, prefer: list[str]) -> str | None:
    keys = [k for k in tracks if k != "live_chat"]
    for k in prefer:
        if k in keys:
            return k
    for k in keys:
        if k.lower().startswith("en"):
            return k
    return None


def _parse_json3(raw: str) -> list[tuple[float, str]]:
    data = json.loads(raw)
    lines = []
    for ev in data.get("events", []):
        segs = ev.get("segs")
        if not segs:
            continue
        text = "".join(s.get("utf8", "") for s in segs).replace("\n", " ").strip()
        if text:
            lines.append((ev.get("tStartMs", 0) / 1000.0, text))
    return lines


def _parse_vtt(raw: str) -> list[tuple[float, str]]:
    def to_sec(ts: str) -> float:
        parts = ts.strip().split(":")
        parts = [float(x.replace(",", ".")) for x in parts]
        while len(parts) < 3:
            parts.insert(0, 0.0)
        return parts[0] * 3600 + parts[1] * 60 + parts[2]

    lines: list[tuple[float, str]] = []
    cur = 0.0
    for line in raw.splitlines():
        s = line.strip()
        if "-->" in s:
            try:
                cur = to_sec(s.split("-->")[0].split()[0])
            except Exception:
                pass
            continue
        if not s or s.startswith(("WEBVTT", "Kind:", "Language:", "NOTE", "STYLE")) or s.isdigit():
            continue
        text = html.unescape(re.sub(r"<[^>]+>", "", s)).strip()
        if text:
            lines.append((cur, text))
    return lines


def _dedupe(lines: list[tuple[float, str]]) -> list[tuple[float, str]]:
    """Drop rolling-caption repeats (a line repeated in the next cue or two)."""
    out: list[tuple[float, str]] = []
    for t, text in lines:
        recent = [x for _, x in out[-3:]]
        if text in recent:
            continue
        # Auto-caption json3 sometimes re-emits the previous line as a prefix.
        if out and text.startswith(out[-1][1]) and len(text) > len(out[-1][1]):
            text = text[len(out[-1][1]):].strip()
            if not text:
                continue
        out.append((t, text))
    return out


def _youtube_transcript_impl(url: str) -> str:
    url = _normalize_yt_url(url)
    vid = _video_id(url)
    ydir = _data_dir() / "youtube"
    ydir.mkdir(parents=True, exist_ok=True)
    cached = ydir / f"{vid}.captions.txt"
    if cached.exists():
        return _truncate(cached.read_text())

    with _ydl(skip_download=True) as ydl:
        info = ydl.extract_info(url, download=False)
        title = info.get("title", "?")
        duration = info.get("duration")
        manual = info.get("subtitles") or {}
        auto = info.get("automatic_captions") or {}

        kind, lang = "manual", _pick_track(manual, ["en", "en-US", "en-GB", "en-CA", "en-AU"])
        tracks = manual
        if lang is None:
            kind, tracks = "auto", auto
            lang = _pick_track(auto, ["en-orig", "en", "en-US"])
        if lang is None:  # non-English manual subs
            other = [k for k in manual if k != "live_chat"]
            if other:
                kind, tracks, lang = "manual", manual, other[0]
        if lang is None:
            orig = [k for k in auto if k.endswith("-orig")]
            if orig:
                kind, tracks, lang = "auto", auto, orig[0]
        if lang is None:
            return (
                f"No captions available for '{title}' ({url}). Try youtube_audio_transcript(url) for "
                "the speech, or youtube_frames_question(url, question) for visual content."
            )

        formats = tracks[lang]
        chosen = None
        for ext in ("json3", "vtt", "srv3", "srv1", "ttml"):
            chosen = next((f for f in formats if f.get("ext") == ext), None)
            if chosen:
                break
        chosen = chosen or formats[0]
        raw = ydl.urlopen(chosen["url"]).read().decode("utf-8", errors="replace")

    ext = chosen.get("ext", "txt")
    (ydir / f"{vid}.{lang}.{kind}.{ext}").write_text(raw)
    if ext == "json3":
        lines = _parse_json3(raw)
    elif ext == "vtt":
        lines = _parse_vtt(raw)
    else:
        lines = [(0.0, html.unescape(re.sub(r"<[^>]+>", " ", raw)))]
    lines = _dedupe(lines)
    if not lines:
        return (
            f"Captions for '{title}' were empty. Try youtube_audio_transcript(url) or "
            "youtube_frames_question(url, question)."
        )
    dur = f", duration {_fmt_ts(duration)}" if duration else ""
    header = f"YouTube video: {title} ({url}{dur})\nCaptions: {kind} '{lang}' ({ext})\n"
    if kind == "auto":
        header += "Note: auto-generated captions can mishear names and numbers.\n"
    text = header + "\n".join(f"[{_fmt_ts(t)}] {line}" for t, line in lines)
    cached.write_text(text)
    return _truncate(text)


def _youtube_audio_impl(url: str, timestamps: bool) -> str:
    url = _normalize_yt_url(url)
    vid = _video_id(url)
    ydir = _data_dir() / "youtube"
    ydir.mkdir(parents=True, exist_ok=True)
    model_name = _whisper_model_name()
    seg_cache = _data_dir() / "transcripts" / f"yt_{vid}_{model_name.replace('/', '-')}.segments.json"
    if seg_cache.exists():
        return _transcribe(Path(f"youtube_{vid}"), timestamps=timestamps, cache_key=f"yt_{vid}")
    tmp = Path(tempfile.mkdtemp(prefix=f"aud_{vid}_", dir=ydir))
    try:
        with _ydl(format="bestaudio/best", outtmpl=str(tmp / "%(id)s.%(ext)s")) as ydl:
            ydl.extract_info(url, download=True)
        files = [f for f in tmp.iterdir() if f.is_file() and not f.name.endswith(".part")]
        if not files:
            return f"ERROR: yt-dlp downloaded no audio for {url}."
        return _transcribe(files[0], timestamps=timestamps, cache_key=f"yt_{vid}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _probe_duration(p: Path) -> float | None:
    try:
        r = subprocess.run(
            [_ffprobe(), "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(p)],
            capture_output=True, text=True, timeout=60,
        )
        return float(r.stdout.strip())
    except Exception:
        return None


def _youtube_frames_impl(url: str, question: str, num_frames: int) -> str:
    from gaia_agent.config import vision_complete  # lazy: config may be mid-build

    url = _normalize_yt_url(url)
    vid = _video_id(url)
    num_frames = max(1, min(int(num_frames or 16), 48))
    ydir = _data_dir() / "youtube"
    ydir.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix=f"frames_{vid}_", dir=ydir))
    try:
        fmt = "bv*[height<=480][ext=mp4]/bv*[height<=480]/b[height<=480]/wv*/w"
        with _ydl(format=fmt, outtmpl=str(tmp / "video.%(ext)s")) as ydl:
            info = ydl.extract_info(url, download=True)
        title = info.get("title", "?")
        videos = [f for f in tmp.iterdir() if f.is_file() and f.name.startswith("video.") and not f.name.endswith(".part")]
        if not videos:
            return f"ERROR: yt-dlp downloaded no video for {url}."
        video = videos[0]
        duration = info.get("duration") or _probe_duration(video) or 0
        if duration <= 0:
            return f"ERROR: could not determine the duration of {url}."
        times = [duration * (i + 0.5) / num_frames for i in range(num_frames)]
        frames: list[tuple[float, Path]] = []
        for i, t in enumerate(times):
            out = tmp / f"frame_{i:03d}.jpg"
            subprocess.run(
                [_ffmpeg(), "-nostdin", "-loglevel", "error", "-ss", f"{t:.2f}", "-i", str(video),
                 "-frames:v", "1", "-vf", "scale=768:-2", "-q:v", "3", "-y", str(out)],
                capture_output=True, timeout=120,
            )
            if out.exists() and out.stat().st_size > 0:
                frames.append((t, out))
        if not frames:
            return f"ERROR: ffmpeg extracted no frames from {url}."
        try:
            video.unlink()  # free the large file before the slow vision calls
        except OSError:
            pass

        answers = []
        batches = [frames[i:i + 8] for i in range(0, len(frames), 8)]
        for bi, batch in enumerate(batches, 1):
            stamps = ", ".join(_fmt_ts(t) for t, _ in batch)
            prompt = (
                f"These {len(batch)} images are frames sampled in order from the YouTube video "
                f"'{title}', at timestamps {stamps} (batch {bi} of {len(batches)}).\n"
                f"Question: {question}\n"
                "First describe, frame by frame, only what is relevant to the question (count "
                "objects carefully and read any on-screen text exactly). Then give your best answer "
                "based on these frames alone, and say how confident you are."
            )
            try:
                ans = vision_complete(prompt, [str(p) for _, p in batch])
            except Exception as e:
                ans = f"ERROR from vision model: {type(e).__name__}: {e}"
            answers.append(f"=== batch {bi} (frames at {stamps}) ===\n{ans}")
        note = (
            f"\n\nCombined note: {len(frames)} frames over {_fmt_ts(duration)} were split into "
            f"{len(batches)} batch(es); each batch saw only its own frames. Combine the batch answers "
            "yourself (for 'the most X at once' take the maximum across batches, not the sum). "
            "Frames are samples, so brief moments can be missed; re-run with more frames if needed."
        )
        return _truncate(f"Video: {title} ({url})\n\n" + "\n\n".join(answers) + note)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _err(e: Exception) -> str:
    return f"ERROR: {type(e).__name__}: {e}"


# ------------------------------------------------------------------------- tools


@tool
def read_file(path: str) -> str:
    """Reads a local file and returns its content as text. Spreadsheets (.xlsx, .xls) come back as
    one markdown table per sheet, with column dtypes, row count, category-like columns, numeric
    column sums and any cell background colours. CSV/TSV likewise. Text and code files (.py, .txt,
    .md, .json, ...) come back as text; PDF, .docx and .pptx have their text extracted; .zip files
    are extracted and their file list returned. For images it tells you to use image_question,
    for audio to use transcribe_audio. Long output is truncated.

    Args:
        path: Local path of the file, e.g. the path returned when the task's attached file was downloaded.
    """
    try:
        return _truncate(_read_file_impl(path))
    except Exception as e:
        return _err(e)


@tool
def run_python_file(path: str, timeout_s: int = 30) -> str:
    """EXECUTES a local Python (.py) file in a separate subprocess with the project's Python
    interpreter and returns its exit code, stdout and stderr. Use it to answer questions like
    "what is the final numeric output of this code". The file really runs, so only use it on the
    task's attached code file.

    Args:
        path: Local path of the .py file to execute.
        timeout_s: Seconds to wait before killing the process (default 30, max 600).
    """
    try:
        p = _resolve(path)
        if p is None:
            return _not_found(path)
        timeout = max(1, min(int(timeout_s or 30), 600))
        try:
            r = subprocess.run(
                [_python_exe(), str(p.resolve())],
                cwd=str(p.resolve().parent),
                capture_output=True,
                text=True,
                timeout=timeout,
                stdin=subprocess.DEVNULL,
            )
        except subprocess.TimeoutExpired as e:
            out = e.stdout.decode(errors="replace") if isinstance(e.stdout, bytes) else (e.stdout or "")
            return _truncate(
                f"ERROR: timed out after {timeout}s (the program may loop for a long time or wait "
                f"for input). Partial stdout:\n{out}"
            )
        text = f"exit code: {r.returncode}\n--- stdout ---\n{r.stdout}"
        if r.stderr.strip():
            text += f"\n--- stderr ---\n{r.stderr}"
        return _truncate(text)
    except Exception as e:
        return _err(e)


@tool
def transcribe_audio(path: str, timestamps: bool = False) -> str:
    """Transcribes speech in a local audio file (.mp3, .wav, .m4a, .flac, ...; video files work
    too) with a local Whisper model and returns the transcript text. Results are cached.

    Args:
        path: Local path of the audio file.
        timestamps: If true, return one line per segment prefixed with [start-end] times.
    """
    try:
        p = _resolve(path)
        if p is None:
            return _not_found(path)
        return _transcribe(p, timestamps=bool(timestamps))
    except Exception as e:
        return _err(e)


@tool
def image_question(path: str, question: str) -> str:
    """Asks a vision model a question about a local image (.png, .jpg, ...) and returns its answer.
    Ask precise, self-contained questions. For a chess board, first ask it to list every piece with
    its square (e.g. "white king g1, black queen d8, ...") and whose turn it is, then reason about
    the position yourself (or check moves with python-chess). For text or tables in an image, ask
    it to transcribe them exactly.

    Args:
        path: Local path of the image file.
        question: The question to ask about the image.
    """
    try:
        p = _resolve(path)
        if p is None:
            return _not_found(path)
        from gaia_agent.config import vision_complete  # lazy: config may be mid-build

        prompt = (
            f"{question}\n\nLook carefully at the whole image. Read any text, numbers or symbols "
            "exactly as shown. If something is unclear, say so instead of guessing."
        )
        return vision_complete(prompt, [str(p)])
    except Exception as e:
        return _err(e)


@tool
def youtube_transcript(url: str) -> str:
    """Returns the captions of a YouTube video as plain text with [mm:ss] timestamps, plus the
    title and duration. Prefers manual English subtitles, then automatic English captions. If the
    video has no captions, it says so: then use youtube_audio_transcript for the speech or
    youtube_frames_question for what is shown on screen.

    Args:
        url: The YouTube video URL (or 11-character video id).
    """
    try:
        return _youtube_transcript_impl(url)
    except Exception as e:
        return _err(e) + " (try youtube_audio_transcript or youtube_frames_question instead)"


@tool
def youtube_audio_transcript(url: str, timestamps: bool = False) -> str:
    """Downloads only the audio of a YouTube video and transcribes it with a local Whisper model.
    Slower than youtube_transcript; use it when the video has no captions or the captions look
    wrong. Results are cached.

    Args:
        url: The YouTube video URL (or 11-character video id).
        timestamps: If true, return one line per segment prefixed with [start-end] times.
    """
    try:
        return _youtube_audio_impl(url, bool(timestamps))
    except Exception as e:
        return _err(e)


@tool
def youtube_frames_question(url: str, question: str, num_frames: int = 16) -> str:
    """Answers a question about what is SHOWN in a YouTube video: downloads a low-resolution copy,
    samples frames evenly across the whole video, and asks a vision model the question in batches
    of up to 8 frames. Returns each batch's answer plus a note on combining them. Use for visual
    questions (counting things on screen, reading on-screen text); use youtube_transcript for
    what is said.

    Args:
        url: The YouTube video URL (or 11-character video id).
        question: The question about the video's visual content.
        num_frames: How many frames to sample evenly across the video (default 16, max 48).
    """
    try:
        return _youtube_frames_impl(url, question, num_frames)
    except Exception as e:
        return _err(e)


TOOLS_MEDIA = [
    read_file,
    run_python_file,
    transcribe_audio,
    image_question,
    youtube_transcript,
    youtube_audio_transcript,
    youtube_frames_question,
]

__all__ = ["TOOLS_MEDIA"] + [t.name for t in TOOLS_MEDIA]
